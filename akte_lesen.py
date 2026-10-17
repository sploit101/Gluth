#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
akte_lesen.py — erstellt die Adresskarte einer Stationsakte.

Was das ist
-----------
Bevor ein Programm in eine PowerPoint schreiben darf, muss es wissen, WOHIN.
Diese Datei geht eine konkrete .pptx durch und haelt zu jeder beschreibbaren
Stelle ihre Adresse fest: Folie, Form, Platzhalternummer, bei Tabellen Zeile
und Spalte.

Das Ergebnis ist unsichtbar und genau deshalb wichtig. Ein Schreiber ohne
verlaessliche Adresskarte ist die Komponente, die irgendwann in die falsche
Zelle schreibt.

Zusaetzlich prueft das Skript, welche Felder aus zielfelder.csv in DIESER
Datei tatsaechlich auffindbar sind. Der Katalog beschreibt die Vorlage; eine
einzelne Akte benutzt davon nur einen Teil.

Ein technischer Punkt
---------------------
PowerPoint zerlegt Text unsichtbar in "runs". Ein Satz kann als
  ["Die Taktzeit muss ", "max. 32", " s betragen."]
gespeichert sein. Wer nach "max. 32 s" sucht, findet nichts. Deshalb liefert
dieses Skript zu jeder Form den ZUSAMMENGEFUEHRTEN Text und zaehlt die runs
mit - eine hohe run-Zahl ist der Hinweis, dass Suchen dort scheitern wird.

Kein Sprachmodell. Aendert nichts, liest nur.

    python akte_lesen.py
    python akte_lesen.py --akte 261008_Stationakte.pptx
    python akte_lesen.py --akte out/Stationsakte_gefuellt.pptx --lang

Ausgabe
-------
    out/akte_karte.json    maschinenlesbar, Eingabe fuer akte_schreiben.py
    out/akte_karte.txt     lesbare Fassung
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from datetime import datetime
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Emu
except ImportError:                                            # pragma: no cover
    raise SystemExit("python-pptx fehlt:  python -m pip install python-pptx")

HERE = Path(__file__).resolve().parent
OUT = HERE / "out"


def ascii_safe(s) -> str:
    return str(s).encode("ascii", "replace").decode("ascii")


def norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


def text_zusammengefuehrt(tf) -> str:
    """
    Text einer Form, ueber alle Absaetze und runs zusammengefuehrt.
    Absatzgrenzen bleiben als Zeilenumbruch erhalten.
    """
    zeilen = []
    for p in tf.paragraphs:
        zeilen.append("".join(r.text for r in p.runs) or p.text or "")
    return "\n".join(zeilen)


def run_anzahl(tf) -> int:
    return sum(len(p.runs) for p in tf.paragraphs)


def vorlage_finden(angabe: str | None) -> Path | None:
    if angabe:
        p = Path(angabe)
        return p if p.exists() else None
    fest = HERE / "vorlage_stationsakte.pptx"
    if fest.exists():
        return fest
    kandidaten = sorted(HERE.glob("*.pptx"), key=lambda p: p.stat().st_mtime,
                        reverse=True)
    if kandidaten:
        print("Keine Datei angegeben, verwende die neueste PowerPoint:")
        print("  ", ascii_safe(kandidaten[0].name))
        print()
        return kandidaten[0]
    return None


def katalog_laden(pfad: Path) -> dict:
    """(Folienart, feld_id) -> Katalogzeile."""
    if not pfad.exists():
        return {}
    kat = {}
    with pfad.open(encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f, delimiter=";"):
            kat[(r["folienart"], r["feld_id"])] = r
    return kat


def karte_bauen(prs, lang: bool = False) -> dict:
    """Die Adresskarte. Eine Liste von Stellen, jede mit vollem Weg."""
    folien = []
    for nr, s in enumerate(prs.slides, start=1):
        eintrag = {
            "folie": nr,
            "folienart": s.slide_layout.name,
            "formen": [],
            "tabellen": [],
        }
        for sh in s.shapes:
            basis = {
                "form_id": sh.shape_id,
                "form_name": sh.name,
                "typ": str(sh.shape_type),
                "links_cm": round(sh.left / 360000, 1) if sh.left else None,
                "oben_cm": round(sh.top / 360000, 1) if sh.top else None,
            }
            if sh.is_placeholder:
                basis["feld_id"] = f"ph{sh.placeholder_format.idx}"
                basis["ph_typ"] = str(sh.placeholder_format.type)
            if sh.has_text_frame:
                t = text_zusammengefuehrt(sh.text_frame)
                basis["text"] = t if lang else norm(t)[:200]
                basis["leer"] = not norm(t)
                basis["runs"] = run_anzahl(sh.text_frame)
                basis["absaetze"] = len(sh.text_frame.paragraphs)
            if sh.has_table:
                tb = sh.table
                gitter = []
                for ri, row in enumerate(tb.rows):
                    zellen = []
                    for ci, c in enumerate(row.cells):
                        zellen.append({
                            "zeile": ri, "spalte": ci,
                            "text": norm(c.text)[:120],
                            "leer": not norm(c.text),
                        })
                    gitter.append(zellen)
                eintrag["tabellen"].append({
                    "form_id": sh.shape_id,
                    "form_name": sh.name,
                    "zeilen": len(tb.rows),
                    "spalten": len(tb.columns),
                    "kopfzeile": [norm(c.text) for c in tb.rows[0].cells]
                    if len(tb.rows) else [],
                    "gitter": gitter,
                })
                basis["ist_tabelle"] = True
            eintrag["formen"].append(basis)
        folien.append(eintrag)
    return {
        "erzeugt": f"{datetime.now():%Y-%m-%d %H:%M}",
        "folien_gesamt": len(folien),
        "folienarten_im_master": [L.name for L in prs.slide_layouts],
        "folien": folien,
    }


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--akte", default=None,
                    help="die .pptx, deren Adresskarte erstellt wird")
    ap.add_argument("--katalog", default=str(HERE / "zielfelder.csv"))
    ap.add_argument("--aus", default=None, help="Ziel-JSON")
    ap.add_argument("--lang", action="store_true",
                    help="vollstaendige Texte in die Karte schreiben "
                         "(VERTRAULICH, Standard ist auf 200 Zeichen gekuerzt)")
    a = ap.parse_args()

    pfad = vorlage_finden(a.akte)
    if pfad is None:
        print("FEHLER: keine .pptx gefunden. --akte angeben.")
        return 2

    prs = Presentation(str(pfad))
    karte = karte_bauen(prs, a.lang)
    karte["quelle"] = pfad.name

    kat = katalog_laden(Path(a.katalog))

    # --------------------------------------------------- Abgleich mit Katalog
    # Welche Katalogfelder sind in DIESER Datei auffindbar?
    gefunden, fehlend = [], []
    vorhanden = set()
    for f in karte["folien"]:
        for sh in f["formen"]:
            if "feld_id" in sh:
                vorhanden.add((f["folienart"], sh["feld_id"]))
    for (fa, fid), zeile in sorted(kat.items()):
        if not fid.startswith("ph"):
            continue                        # Tabellen- und Grafikfelder separat
        if (fa, fid) in vorhanden:
            gefunden.append((fa, fid, zeile["feldname"], zeile["befuellbar"]))
        else:
            fehlend.append((fa, fid, zeile["feldname"], zeile["befuellbar"]))

    tabellen = [(f["folie"], t["form_name"], t["zeilen"], t["spalten"],
                 t["kopfzeile"])
                for f in karte["folien"] for t in f["tabellen"]]

    zerstueckelt = [(f["folie"], sh["form_name"], sh["runs"])
                    for f in karte["folien"] for sh in f["formen"]
                    if sh.get("runs", 0) > 3]

    # ------------------------------------------------------------- Ausgabe
    L = [f"ADRESSKARTE  {karte['quelle']}  {karte['erzeugt']}", "=" * 70, ""]
    L += [f"  Folien in der Datei             {karte['folien_gesamt']}",
          f"  Folienarten im Master           {len(karte['folienarten_im_master'])}",
          f"  Katalogfelder auffindbar        {len(gefunden)}",
          f"  Katalogfelder nicht in dieser Datei  {len(fehlend)}",
          f"  Tabellen gefunden               {len(tabellen)}",
          f"  Formen mit zerstueckeltem Text  {len(zerstueckelt)}",
          "", "  Folien:"]
    for f in karte["folien"]:
        felder = [sh["feld_id"] for sh in f["formen"] if "feld_id" in sh]
        leer = sum(1 for sh in f["formen"] if sh.get("leer"))
        L.append(f"    {f['folie']:>2}  {f['folienart'][:30]:<30} "
                 f"{len(f['formen'])} Formen, {len(felder)} Platzhalter, "
                 f"{leer} leer")
    if tabellen:
        L += ["", "  Tabellen:"]
        for nr, name, z, s, kopf in tabellen:
            L.append(f"    Folie {nr}  {name}  {z} x {s}")
            L.append(f"              Kopfzeile: {' | '.join(kopf)[:90]}")
    if zerstueckelt:
        L += ["", "  Zerstueckelter Text (Suchen wuerde dort fehlschlagen):"]
        for nr, name, n in zerstueckelt[:10]:
            L.append(f"    Folie {nr}  {name}  {n} Textteile")
        L.append("    -> akte_schreiben.py fuehrt vor dem Suchen zusammen.")

    L += ["", "  Auffindbare Katalogfelder, die befuellbar sind:"]
    n = 0
    for fa, fid, name, bf in gefunden:
        if bf in ("ja", "teilweise"):
            L.append(f"    {fa[:28]:<28} {fid:<6} {name[:34]:<34} {bf}")
            n += 1
    if not n:
        L.append("    keine")

    L += ["", "  Lesart:",
          "    'nicht in dieser Datei' ist normal. Der Katalog beschreibt die",
          "    VORLAGE mit allen 21 Folienarten; eine einzelne Akte benutzt",
          "    davon nur einige. Fehlende Folien legt der Schreiber bei Bedarf",
          "    aus der passenden Folienart neu an.",
          "", "Diese Datei enthaelt gekuerzte Texte, aber Dateinamen."]

    OUT.mkdir(exist_ok=True)
    ziel_txt = OUT / "akte_karte.txt"
    ziel_txt.write_text("\n".join(L), encoding="utf-8")
    for z in L:
        print(ascii_safe(z))

    ziel = Path(a.aus) if a.aus else OUT / "akte_karte.json"
    karte["abgleich"] = {
        "gefunden": [{"folienart": fa, "feld_id": fid, "feldname": nm,
                      "befuellbar": bf} for fa, fid, nm, bf in gefunden],
        "fehlend": [{"folienart": fa, "feld_id": fid, "feldname": nm,
                     "befuellbar": bf} for fa, fid, nm, bf in fehlend],
    }
    ziel.write_text(json.dumps(karte, ensure_ascii=False, indent=1),
                    encoding="utf-8")

    print()
    print("geschrieben:")
    print(f"  {ziel}")
    print(f"  {ziel_txt}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
