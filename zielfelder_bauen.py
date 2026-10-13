#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
zielfelder_bauen.py — erzeugt den Zielfeld-Katalog aus der Stationsakte-Vorlage.

Hintergrund
-----------
Die Vorlage enthaelt in den Folienlayouts Platzhalter mit HINWEISTEXT, den
Gluth selbst geschrieben hat ("Zeichnungsnr.", "Auszug aus Lastenheft/Angebot",
"Stunden", "U/min"). Das ist faktisch ein Feldverzeichnis aus erster Hand.
Dieses Skript liest es aus und schreibt daraus eine Tabelle, in der fuer jedes
beschreibbare Feld steht, WOHER sein Inhalt kommen kann.

Diese Tabelle ist die Spezifikation der Erstbefuellung. Der Zuordner liest sie
spaeter ein; vorher wird sie von einem Menschen korrigiert.

Spalten
-------
nr            laufende Nummer
folienart     Name des Layouts in der Vorlage
feld_id       Adresse des Feldes: "ph<idx>" oder "tab:<Tabelle>:<Spalte>"
feldname      Hinweistext aus der Vorlage, WOERTLICH uebernommen
typ           Text | Zahl+Einheit | Datum | Name | Ja/Nein | Bild | Liste
quelle        Lastenheft | Spezifikation | Protokoll | Konstruktion |
              Einkauf | Projektplanung | Automatisch | Bild | unklar
befuellbar    ja | teilweise | nein   (automatisch befuellbar)
sicherheit    hoch | mittel | niedrig (wie sicher ist unsere Einschaetzung)
prio          1 zuerst bauen, 2 spaeter, 3 nicht automatisierbar
frage         offene Frage an Daniel Genau, sonst leer
bemerkung     Begruendung oder Hinweis

Aufruf
------
    python zielfelder_bauen.py
    python zielfelder_bauen.py --vorlage vorlage_stationsakte.pptx --aus zielfelder.csv

Nach jeder neuen Fassung der Vorlage neu laufen lassen: die Struktur kommt dann
aus der Datei, nur die Einschaetzungsspalten muessen nachgezogen werden.
"""

from __future__ import annotations

import argparse
import collections
import csv
import re
from pathlib import Path

try:
    from pptx import Presentation
except ImportError:                                            # pragma: no cover
    raise SystemExit("python-pptx fehlt:  python -m pip install python-pptx")

HERE = Path(__file__).resolve().parent

# ---------------------------------------------------------------------------
# Einschaetzung je Layout und Platzhalter-Index.
#
# Tupel: (typ, quelle, befuellbar, sicherheit, prio, frage, bemerkung)
#
# "*" als Index gilt fuer alle Platzhalter des Layouts, fuer die kein
# eigener Eintrag existiert. So bleibt die Tabelle kurz und lesbar.
# ---------------------------------------------------------------------------

BILD = ("Bild", "Bild", "teilweise", "mittel", 3, "",
        "Bilder aus dem PDF extrahierbar, aber die Auswahl des richtigen "
        "Bildes ist nicht verlaesslich automatisierbar. Bleibt leer und "
        "wird als Luecke markiert.")

KONSTRUKTION = ("Zahl+Einheit", "Konstruktion", "nein", "hoch", 3, "",
                "Auslegungsergebnis. Steht in keinem Kundendokument.")

ZEIT = ("Zahl+Einheit", "Konstruktion", "teilweise", "niedrig", 2,
        "Ist mit 'Zeit' die Taktzeit gemeint? Die waere Kundenvorgabe.",
        "Grenzfall: Taktzeiten gibt der Kunde vor, Einzelzeiten rechnet "
        "die Konstruktion.")

EINSCHAETZUNG: dict[str, dict] = {

    # Achtung: der Layoutname ist in der Vorlage mit Schreibfehler angelegt
    # ("Stationsbeschriebung"). Hier steht die Schreibweise AUS DER DATEI,
    # sonst greift die Einschaetzung nicht.
    "Stationsbeschriebung": {
        0:  ("Text", "Lastenheft", "teilweise", "niedrig", 2,
             "Gibt es ein festes Schema fuer Stationsnummern und "
             "-bezeichnungen, und kommt die Nummer von euch oder vom Kunden?",
             "Bezeichnung oft aus dem Lastenheft, Nummer vermutlich intern."),
        11: ("Text", "Lastenheft", "ja", "hoch", 1, "",
             "Beschreibung der Stationsinhalte - woertliche Auszuege moeglich."),
        13: ("Zahl+Einheit", "Projektplanung", "nein", "hoch", 3, "",
             "Stunden. Interne Kalkulation."),
        14: ("Zahl+Einheit", "Projektplanung", "nein", "hoch", 3, "",
             "Budget. Interne Kalkulation."),
        15: ("Zahl+Einheit", "Projektplanung", "nein", "hoch", 3, "",
             "Budget. Interne Kalkulation."),
        16: ("Datum", "Projektplanung", "nein", "hoch", 3, "",
             "Termin. Aus der Projektplanung, nicht aus dem Dokument."),
        17: ("Datum", "Projektplanung", "nein", "hoch", 3, "",
             "Termin. Aus der Projektplanung, nicht aus dem Dokument."),
        "*": BILD,
    },

    "Überschrift": {
        0:  ("Text", "Automatisch", "ja", "hoch", 1, "",
             "Reine Abschnittsueberschrift. Setzt das Programm selbst."),
    },

    "K-Prozess": {
        0:  ("Text", "unklar", "teilweise", "niedrig", 2,
             "Wofuer ist die Folienart 'K-Prozess' gedacht? In der Vorlage "
             "hat sie nur einen Titel und keine weiteren Felder.",
             "Nicht einzuordnen, solange die Bedeutung offen ist. "
             "Vermutung: Kundenprozess oder Kritischer Prozess."),
    },

    "Leer": {},

    "Aufgabenstellung_ausführlich": {
        11: ("Text", "Lastenheft", "ja", "hoch", 1, "",
             "STAERKSTES FELD. Der Hinweistext nennt ausdruecklich "
             "'Auszug aus Lastenheft/Angebot' - genau das liefert die "
             "Extraktion heute schon woertlich mit Datei und Seite."),
        "*": BILD,
    },

    "technische Ausführung": {
        11: ("Text", "Konstruktion", "nein", "hoch", 3, "",
             "Problemstellung, Loesungsfindung, Loesungsansatz. "
             "Das ist die Denkleistung der Konstruktion."),
        "*": BILD,
    },

    "Kalibrier-/Einrichtteile": {
        11: ("Text", "Spezifikation", "teilweise", "mittel", 2,
             "Werden Kalibrier- und Einrichtablaeufe vom Kunden "
             "vorgegeben oder von euch festgelegt?",
             "Im Beispiel der Vorlage stehen auf der Folie 'Sonderablaeufe' "
             "Angaben wie 'CTC Ablauf nach Schema xy' - das klingt nach "
             "Kundenvorgabe und waere damit befuellbar."),
        "*": BILD,
    },

    "Auslegung Kaufteil": {
        11: ("Text", "Lastenheft", "teilweise", "niedrig", 2, "",
             "Der Hinweistext nennt 'Anforderungen, Kraefte, Druecke, "
             "Schnittstellen'. Die Anforderungen koennen zitiert werden, "
             "die Bauteilauswahl nicht."),
        "*": BILD,
    },

    "Auslegung Feder": {
        11: ("Text", "Konstruktion", "nein", "mittel", 3, "", "Freitext 'Sonstiges'."),
        13: KONSTRUKTION, 14: KONSTRUKTION, 15: KONSTRUKTION, 16: KONSTRUKTION,
        "*": BILD,
    },

    "Auslegung Roboter": {
        11: ("Text", "Konstruktion", "nein", "mittel", 3, "", "Freitext 'Sonstiges'."),
        16: ZEIT,
        "*": KONSTRUKTION,
    },

    "Auslegung Achssystem": {
        11: ("Text", "Konstruktion", "nein", "mittel", 3, "", "Freitext 'Sonstiges'."),
        16: ZEIT,
        "*": KONSTRUKTION,
    },

    "Auslegung Antrieb/Getriebe": {
        11: ("Text", "Konstruktion", "nein", "mittel", 3, "", "Freitext 'Sonstiges'."),
        16: ZEIT,
        "*": KONSTRUKTION,
    },

    "Auslegung Bildverarbeitung": {
        11: ("Text", "Konstruktion", "nein", "mittel", 3, "", "Freitext 'Sonstiges'."),
        16: ZEIT,
        "*": KONSTRUKTION,
    },

    "Auslegung Absaugung": {
        11: ("Text", "Konstruktion", "nein", "mittel", 3, "", "Prozess, Absaugpunkte."),
        "*": KONSTRUKTION,
    },

    "Auslegung Kabelschlepp": {
        11: ("Text", "Konstruktion", "nein", "mittel", 3, "",
             "Prozess, Bewegungsprofil."),
        "*": KONSTRUKTION,
    },

    "Bauteile": {
        # Das Layout enthaelt ZWEI Bauteilbloecke: idx 13-22 und idx 24-33.
        # Beide fragen dasselbe ab, deshalb wird die Einschaetzung unten
        # programmatisch gespiegelt (siehe BAUTEILE_BLOCK2).
        0:  ("Text", "Lastenheft", "ja", "hoch", 1, "", "Ueberschrift Kundenbauteile."),
        13: ("Text", "Lastenheft", "ja", "hoch", 1, "", "Bauteilname (Bauteil 1)."),
        14: ("Text", "Lastenheft", "ja", "hoch", 1, "",
             "Zeichnungsnummer (Bauteil 1). Gut per Muster erkennbar."),
        15: ("Text", "Lastenheft", "ja", "mittel", 1, "", "Nr./Index (Bauteil 1)."),
        16: ("Ja/Nein", "unklar", "teilweise", "niedrig", 2,
             "Was wird in den beiden 'Ja/nein'-Feldern der Folienart "
             "'Bauteile' abgefragt?",
             "Bedeutung nicht aus der Vorlage erkennbar."),
        17: ("Ja/Nein", "unklar", "teilweise", "niedrig", 2, "",
             "Bedeutung nicht aus der Vorlage erkennbar."),
        18: ("Liste", "Lastenheft", "teilweise", "mittel", 2, "",
             "Anlieferform (Schuettgut/Blister). Steht manchmal im Lastenheft."),
        19: ("Text", "Lastenheft", "teilweise", "mittel", 2, "", "Text oder Skizze."),
        20: ("Text", "Lastenheft", "ja", "mittel", 1, "", "Nr./Index (Bauteil 1)."),
        21: ("Text", "Lastenheft", "teilweise", "mittel", 2, "",
             "Hinweise zu Produktaenderungen und Variantenrahmen. "
             "Steht oft im Lastenheft, aber verstreut."),
        22: ("Ja/Nein", "Lastenheft", "ja", "mittel", 1, "", "Existieren Varianten."),
    },

    "Kundenänderung": {
        0:  ("Text", "Lastenheft", "ja", "mittel", 2, "",
             "Bezeichnung der Aenderung. Quelle ist die neue Fassung des "
             "Lastenhefts - das ist der zweite Schritt, die Updates."),
        11: ("Text", "Lastenheft", "ja", "mittel", 2, "", "Beschreibung der Aufgabe."),
        21: ("Text", "Lastenheft", "teilweise", "mittel", 2, "",
             "Was fuehrte zum Aenderungswunsch. Der Versionsvergleich der "
             "Lastenhefte liefert genau das: alter Wert gegen neuer Wert."),
        "*": ("Zahl+Einheit", "Projektplanung", "nein", "hoch", 3, "",
              "Stunden und Budget. Interne Kalkulation."),
    },

    "Interne Änderung": {
        0:  ("Text", "Protokoll", "teilweise", "niedrig", 2, "",
             "Interne Aenderung. Quelle waere das Besprechungsprotokoll."),
        11: ("Text", "Protokoll", "teilweise", "niedrig", 2, "",
             "Beschreibung der Aufgabe."),
        21: ("Text", "Protokoll", "teilweise", "niedrig", 2, "",
             "Was fuehrte zum Aenderungswunsch."),
        22: ("Zahl+Einheit", "Projektplanung", "nein", "hoch", 3, "", "Kosten."),
        "*": ("Zahl+Einheit", "Projektplanung", "nein", "hoch", 3, "",
              "Stunden und Budget. Interne Kalkulation."),
    },

    "Kaufteil neu": {
        0:  ("Text", "Einkauf", "nein", "mittel", 3, "", "Bezeichnung."),
        11: ("Text", "Konstruktion", "nein", "hoch", 3, "",
             "Begruendung der Lieferanten- oder Bauteilwahl. "
             "Nachvollziehbarkeitstext, rein intern."),
        13: ("Zahl+Einheit", "Einkauf", "nein", "hoch", 3, "", "Preis."),
        14: ("Zahl+Einheit", "Einkauf", "nein", "hoch", 3, "", "Lieferzeit in Wochen."),
        "*": ("Name", "Einkauf", "nein", "hoch", 3, "", "Lieferantenname."),
    },

    "Konstruktionsfreigabe": {
        0:  ("Text", "Protokoll", "teilweise", "mittel", 2, "",
             "Stationsnummer und Bezeichnung."),
        11: ("Text", "Protokoll", "ja", "mittel", 2, "",
             "Mitschrieb der Besprechung. Quelle ist das Protokoll - "
             "Schritt 3. Format noch unbekannt."),
        13: ("Datum", "Protokoll", "ja", "hoch", 2, "", "Datum der Besprechung."),
        19: ("Name", "Protokoll", "ja", "mittel", 2, "", "Teilnehmerliste."),
        26: ("Datum", "Protokoll", "ja", "hoch", 2, "", "Datum."),
        "*": ("Name", "Protokoll", "nein", "hoch", 3, "",
              "Unterschrift. Bleibt zwingend manuell."),
    },

    "Besprechungsprotokoll": {
        0:  ("Text", "Protokoll", "teilweise", "mittel", 2, "",
             "Stationsnummer und Bezeichnung."),
        11: ("Text", "Protokoll", "ja", "mittel", 2, "",
             "Mitschrieb und gemeinsame Festlegungen. Kernfeld fuer Schritt 3."),
        13: ("Datum", "Protokoll", "ja", "hoch", 2, "", "Datum der Besprechung."),
        19: ("Name", "Protokoll", "ja", "mittel", 2, "", "Teilnehmerliste."),
        "*": ("Name", "Protokoll", "nein", "hoch", 3, "",
              "Unterschrift. Bleibt zwingend manuell."),
    },
}

STANDARD = ("Text", "unklar", "teilweise", "niedrig", 2, "",
            "Noch nicht eingeschaetzt.")

# Zweiter Bauteilblock (idx 24-33) spiegelt den ersten (idx 13-22).
BAUTEILE_BLOCK2 = {
    neu: EINSCHAETZUNG["Bauteile"][alt][:5] + (
        "",                                        # Frage nicht doppelt stellen
        EINSCHAETZUNG["Bauteile"][alt][6].replace("Bauteil 1", "Bauteil 2"),)
    for alt, neu in zip(range(13, 23), range(24, 34))
}
EINSCHAETZUNG["Bauteile"].update(BAUTEILE_BLOCK2)

# ---------------------------------------------------------------------------
# Felder, die keine Platzhalter sind: Tabellenspalten und freie Textfelder
# aus den Beispielfolien. Die muessen von Hand gepflegt werden, weil sie in
# den Layouts nicht vorkommen.
# ---------------------------------------------------------------------------

TABELLEN_FELDER = [
    # (folienart, feld_id, feldname, typ, quelle, befuellbar, sicherheit,
    #  prio, frage, bemerkung)
    ("Versionsübersicht (Tabelle)", "tab:Version:Datum", "Datum",
     "Datum", "Automatisch", "ja", "hoch", 1, "",
     "Das Programm schreibt bei jeder Befuellung eine Zeile. Freier Gewinn."),
    ("Versionsübersicht (Tabelle)", "tab:Version:Version", "Version",
     "Text", "Automatisch", "ja", "hoch", 1, "",
     "Fortlaufend durch das Programm."),
    ("Versionsübersicht (Tabelle)", "tab:Version:Änderungen", "Änderungen",
     "Text", "Automatisch", "ja", "hoch", 1, "",
     "Kurzfassung aus dem Aenderungsprotokoll des Schreibers."),
    ("Versionsübersicht (Tabelle)", "tab:Version:Verantwortlich", "Verantwortlich",
     "Name", "Automatisch", "ja", "mittel", 1,
     "Soll in der Versionsuebersicht bei automatischer Befuellung der "
     "Name des Programms oder der Name des Freigebenden stehen?",
     "Vorschlag: Name des Menschen, der freigegeben hat, plus Hinweis "
     "'automatisch vorbefuellt'."),

    ("Fähigkeitstabelle", "tab:Faehigkeit:Prozess", "Prozess",
     "Text", "Spezifikation", "ja", "hoch", 1, "",
     "Im Beispiel 'Messung Lagersitztiefe'. Prozessbezeichnung aus der "
     "Spezifikation."),
    ("Fähigkeitstabelle", "tab:Faehigkeit:Verfahren", "Verfahren",
     "Text", "Spezifikation", "ja", "hoch", 1, "",
     "Im Beispiel 'MSA1', 'MSA3'. Kurzkennungen, gut erkennbar."),
    ("Fähigkeitstabelle", "tab:Faehigkeit:Kennwert", "Cmk/cpk/grr",
     "Text", "Spezifikation", "ja", "hoch", 1, "",
     "Im Beispiel 'Cg/cgk', 'Grr'. Feste Begriffsliste."),
    ("Fähigkeitstabelle", "tab:Faehigkeit:Toleranz", "Toleranz",
     "Zahl+Einheit", "Spezifikation", "ja", "hoch", 1, "",
     "Im Beispiel '30 Mikrometer', '<=10 %'. Operator, Wert und Einheit "
     "liest requirements.py heute schon aus."),
    ("Fähigkeitstabelle", "tab:Faehigkeit:Kommentar", "Kommentar",
     "Text", "Konstruktion", "nein", "mittel", 3, "",
     "Freifeld fuer den Bearbeiter. Bleibt leer."),
    ("Fähigkeitstabelle", "tab:Faehigkeit:Kennung", "Nummerierung der Zeilen",
     "Text", "Automatisch", "ja", "hoch", 1, "",
     "Die vorgeschlagene Kennung R-0001. Erzeugt das Programm bei der "
     "Erstbefuellung. Spalte ist in der Vorlage als Kopf schon angelegt."),

    ("Bauteile/Varianten (Grafik)", "txt:Variante", "Kundenbauteil Typ A..D",
     "Text", "Lastenheft", "teilweise", "mittel", 2, "",
     "Textfelder neben den Kreisen. Variantennamen aus dem Lastenheft "
     "moeglich, die Bilder in den Kreisen nicht."),
]


def ascii_safe(s) -> str:
    """Fuer die Windows-Konsole: Umlaute ersetzen statt abbrechen."""
    return str(s).encode("ascii", "replace").decode("ascii")


def saeubern(s: str) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


def einschaetzen(layout: str, idx, ph_typ: str = "") -> tuple:
    """
    Einschaetzung fuer ein Zielfeld.

    Bildplatzhalter werden IMMER als Bild eingeordnet, noch vor jeder
    Layoutregel. Der Platzhaltertyp steht in der Datei und ist damit
    Tatsache, keine Vermutung. (Fehler der ersten Fassung: die Sammelregel
    eines Layouts hat Bildplatzhalter ueberschrieben, dadurch stand bei
    einem Screenshot-Feld "Unterschrift".)
    """
    if "PICTURE" in str(ph_typ).upper():
        return BILD
    regeln = EINSCHAETZUNG.get(layout)
    if regeln is None:
        return STANDARD
    if idx in regeln:
        return regeln[idx]
    if "*" in regeln:
        return regeln["*"]
    return STANDARD


def typ_raten(prompt: str, vorgabe: str) -> str:
    """Nur fuer Sammelregeln: Typ anhand des Hinweistextes nachjustieren."""
    p = prompt.lower()
    if re.search(r"hersteller|name|typ\b", p):
        return "Name"
    if re.search(r"ja/nein|ja/nein/", p):
        return "Ja/Nein"
    if re.search(r"datum|termin", p):
        return "Datum"
    if re.search(r"bild|skizze|foto|screenshot|layout", p):
        return "Bild"
    if re.search(r"anzahl|weg|zeit|gewicht|ma[sß]|radius|breite|h[oö]he|"
                 r"u/min|kgm|m³|µ|mm|preis|stunden|budget|kraft|l[aä]nge|"
                 r"pixel|toleranz|wochen", p):
        return "Zahl+Einheit"
    return vorgabe


KORREKTUR_SPALTEN = ["folienart", "feld_id", "quelle", "befuellbar",
                     "sicherheit", "prio", "bemerkung"]


def vorlage_finden(angabe: str | None) -> Path | None:
    """
    Vorlage bestimmen. Reihenfolge:
      1. ausdruecklich angegebener Pfad
      2. vorlage_stationsakte.pptx im Projektordner
      3. die neueste .pptx im Projektordner  (dann wird sie genannt)
    Damit entfaellt die Dateikopie, die es vorher brauchte.
    """
    if angabe:
        p = Path(angabe)
        return p if p.exists() else None
    fest = HERE / "vorlage_stationsakte.pptx"
    if fest.exists():
        return fest
    kandidaten = sorted(HERE.glob("*.pptx"), key=lambda p: p.stat().st_mtime,
                        reverse=True)
    if kandidaten:
        print("Keine vorlage_stationsakte.pptx vorhanden, verwende die")
        print("neueste PowerPoint im Ordner:", ascii_safe(kandidaten[0].name))
        print()
        return kandidaten[0]
    return None


def korrekturen_laden(pfad: Path) -> dict[tuple[str, str], dict]:
    """Manuelle Korrekturen einlesen. Leere Zellen aendern nichts."""
    if not pfad.exists():
        return {}
    raus: dict[tuple[str, str], dict] = {}
    with pfad.open(encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f, delimiter=";"):
            fa = (r.get("folienart") or "").strip()
            fid = (r.get("feld_id") or "").strip()
            if not fa or fa.startswith("#"):
                continue
            aend = {k: (r.get(k) or "").strip()
                    for k in KORREKTUR_SPALTEN[2:]
                    if (r.get(k) or "").strip()}
            if aend:
                raus[(fa, fid)] = aend
    return raus


def korrekturvorlage_schreiben(pfad: Path) -> None:
    """Leere Korrekturdatei anlegen, damit es etwas zum Ausfuellen gibt."""
    with pfad.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(KORREKTUR_SPALTEN)
        w.writerow(["# Beispielzeile, wird ignoriert. Zum Korrigieren:",
                    "folienart und feld_id aus zielfelder.csv abschreiben,",
                    "dann nur die Spalten ausfuellen, die sich aendern sollen.",
                    "", "", "", ""])
        w.writerow(["# Bauteile", "ph14", "", "nein", "hoch", "3",
                    "Zeichnungsnummern kommen im Lastenheft nicht vor"])


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--vorlage", default=None,
                    help="Pfad zur Stationsakte-Vorlage (.pptx)")
    ap.add_argument("--aus", default=str(HERE / "zielfelder.csv"))
    ap.add_argument("--korrekturen",
                    default=str(HERE / "zielfelder_korrekturen.csv"),
                    help="eigene Korrekturen, werden zuletzt angewendet")
    a = ap.parse_args()

    pfad = vorlage_finden(a.vorlage)
    if pfad is None:
        print("FEHLER: keine Vorlage gefunden.")
        print("Eine .pptx in den Projektordner legen oder --vorlage angeben.")
        return 2

    prs = Presentation(str(pfad))

    # Welche Layouts kommen in der Beispieldatei tatsaechlich vor?
    genutzt = collections.Counter(s.slide_layout.name for s in prs.slides)

    zeilen = []
    for layout in prs.slide_layouts:
        phs = sorted(layout.placeholders,
                     key=lambda sh: sh.placeholder_format.idx)
        if not phs:
            zeilen.append({
                "folienart": layout.name, "feld_id": "-",
                "feldname": "(keine Platzhalter)", "typ": "-",
                "quelle": "-", "befuellbar": "nein", "sicherheit": "hoch",
                "prio": 3, "frage": "",
                "bemerkung": "Leeres Layout, kein Zielfeld.",
                "im_beispiel": genutzt.get(layout.name, 0),
            })
            continue
        for sh in phs:
            idx = sh.placeholder_format.idx
            ph_typ = str(sh.placeholder_format.type)
            prompt = saeubern(sh.text_frame.text) if sh.has_text_frame else ""
            typ, quelle, befuellbar, sicherheit, prio, frage, bem = \
                einschaetzen(layout.name, idx, ph_typ)
            # Bei Sammelregeln den Typ am Hinweistext nachschaerfen
            if (idx not in EINSCHAETZUNG.get(layout.name, {})
                    and "PICTURE" not in ph_typ.upper()):
                typ = typ_raten(prompt, typ)
            zeilen.append({
                "folienart": layout.name,
                "feld_id": f"ph{idx}",
                "feldname": prompt or "(ohne Hinweistext)",
                "typ": typ, "quelle": quelle, "befuellbar": befuellbar,
                "sicherheit": sicherheit, "prio": prio, "frage": frage,
                "bemerkung": bem,
                "im_beispiel": genutzt.get(layout.name, 0),
            })

    for (fa, fid, fn, typ, q, bf, si, pr, fr, bem) in TABELLEN_FELDER:
        zeilen.append({
            "folienart": fa, "feld_id": fid, "feldname": fn, "typ": typ,
            "quelle": q, "befuellbar": bf, "sicherheit": si, "prio": pr,
            "frage": fr, "bemerkung": bem, "im_beispiel": 1,
        })

    # ------------------------------------------------- eigene Korrekturen
    kpfad = Path(a.korrekturen)
    if not kpfad.exists():
        korrekturvorlage_schreiben(kpfad)
        print("Korrekturdatei neu angelegt:", ascii_safe(kpfad.name))
        print()
    korr = korrekturen_laden(kpfad)
    angewendet, unbekannt = 0, []
    if korr:
        schluessel = {(z["folienart"], z["feld_id"]) for z in zeilen}
        for (fa, fid), aend in korr.items():
            if (fa, fid) not in schluessel:
                unbekannt.append(f"{fa} / {fid}")
                continue
            for z in zeilen:
                if z["folienart"] == fa and z["feld_id"] == fid:
                    z.update(aend)
                    z["bemerkung"] = (z.get("bemerkung", "") +
                                      "  [korrigiert von Hand]").strip() \
                        if "bemerkung" not in aend else \
                        aend["bemerkung"] + "  [korrigiert von Hand]"
                    angewendet += 1

    spalten = ["nr", "folienart", "feld_id", "feldname", "typ", "quelle",
               "befuellbar", "sicherheit", "prio", "im_beispiel",
               "frage", "bemerkung"]

    ziel = Path(a.aus)
    with ziel.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=spalten, delimiter=";")
        w.writeheader()
        for n, z in enumerate(zeilen, start=1):
            z["nr"] = n
            w.writerow({k: z.get(k, "") for k in spalten})

    # ------------------------------------------------------------- Zahlen
    print(f"ZIELFELD-KATALOG  aus  {pfad.name}")
    print("=" * 66)
    print(f"  Folienarten in der Vorlage      {len(prs.slide_layouts)}")
    print(f"  davon in der Beispieldatei      {len(genutzt)}")
    print(f"  Zielfelder gesamt               {len(zeilen)}")
    print()
    print("  Automatisch befuellbar")
    bf = collections.Counter(z["befuellbar"] for z in zeilen)
    for k in ("ja", "teilweise", "nein"):
        n = bf.get(k, 0)
        print(f"    {n:>5}  {k:<10} {100*n/len(zeilen):5.1f} %")
    print()
    print("  Quelle des Inhalts")
    for k, n in collections.Counter(z["quelle"] for z in zeilen).most_common():
        print(f"    {n:>5}  {k}")
    print()
    print("  Sicherheit unserer Einschaetzung")
    for k in ("hoch", "mittel", "niedrig"):
        n = sum(1 for z in zeilen if z["sicherheit"] == k)
        print(f"    {n:>5}  {k}")
    print()
    fragen = sorted({z["frage"] for z in zeilen if z["frage"]})
    print(f"  Offene Fragen an Daniel Genau   {len(fragen)}")
    for q in fragen:
        print(f"    - {q}")
    print()
    print(f"  Eigene Korrekturen angewendet   {angewendet}")
    if unbekannt:
        print("  ACHTUNG, diese Korrekturzeilen passen auf kein Feld:")
        for u in unbekannt:
            print("    -", ascii_safe(u))
        print("    folienart und feld_id genau aus zielfelder.csv abschreiben.")
    print()
    print("geschrieben:", ziel.name)
    print("  Spalten quelle/befuellbar/sicherheit sind UNSERE Einschaetzung")
    print("  und gehoeren gegengelesen, nicht als Tatsache behandelt.")
    print(f"  Eigene Aenderungen in {ascii_safe(kpfad.name)} eintragen -")
    print("  die ueberleben jeden neuen Lauf. Aenderungen direkt in")
    print("  zielfelder.csv werden hier ueberschrieben.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
