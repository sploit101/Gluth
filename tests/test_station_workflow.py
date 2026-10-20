"""Synthetic checks for reviewed decisions and large station tables."""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Cm

from akte_schreiben import freigaben_laden, freigabe_schluessel


ROOT = Path(__file__).resolve().parents[1]


class StationWorkflowTests(unittest.TestCase):
    def test_review_must_match_current_mapping(self):
        mapping = {
            "kennung": "R-0001", "folienart": "Station", "feld_id": "ph11",
            "text": "The station shall inspect every part.", "datei": "demo.pdf",
            "seite": "2", "station": "Station 20",
        }
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "review.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.append(["Kennung", "Folienart", "Feld", "Wortlaut", "Datei",
                          "Seite", "Station", "Freigabe"])
            sheet.append([mapping[name] for name in ("kennung", "folienart",
                          "feld_id", "text", "datei", "seite", "station")] + ["ja"])
            book.save(path)

            approvals, error = freigaben_laden(path, [mapping])
            self.assertEqual(error, "")
            self.assertEqual(approvals[freigabe_schluessel(*(mapping[name] for name
                in ("kennung", "folienart", "feld_id", "text", "datei", "seite",
                    "station")))], "ja")

            changed = {**mapping, "station": "Station 30"}
            _, error = freigaben_laden(path, [changed])
            self.assertIn("aktuellen Zuordnung", error)

            sheet.cell(2, 8).value = "maybe"
            book.save(path)
            _, error = freigaben_laden(path, [mapping])
            self.assertIn("Ungueltige Freigabe", error)

    def test_large_table_stays_within_slides(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "out").mkdir()
            shutil.copy2(ROOT / "akte_schreiben.py", root / "akte_schreiben.py")
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[5])
            table = slide.shapes.add_table(2, 6, Cm(1), Cm(4), Cm(23), Cm(2)).table
            for cell, title in zip(table.rows[0].cells,
                                   ("Prozess", "Verfahren", "Cmk/cpk/grr",
                                    "Toleranz", "Kommentar", "Kennung")):
                cell.text = title
            template = root / "template.pptx"
            deck.save(template)
            rows = [
                {"kennung": f"R-{i:04d}", "folienart": slide.slide_layout.name,
                 "feld_id": "tab:Faehigkeit:Prozess", "text": f"Inspect process {i}.",
                 "auszug": "", "pruefen": "", "datei": "synthetic.pdf",
                 "seite": str(i), "station": "Station 20"}
                for i in range(1, 101)
            ]
            mapping = root / "mapping.jsonl"
            mapping.write_text("".join(json.dumps(row) + "\n" for row in rows),
                               encoding="utf-8")
            output = root / "output.pptx"
            run = subprocess.run([
                sys.executable, str(root / "akte_schreiben.py"),
                "--vorlage", str(template), "--zuordnung", str(mapping),
                "--katalog", str(root / "missing.csv"), "--aus", str(output),
                "--station", "Station 20", "--nur-sichere",
            ], cwd=root, capture_output=True, text=True)
            self.assertEqual(run.returncode, 0, run.stdout + run.stderr)

            result = Presentation(output)
            table_slides = 0
            written = 0
            for result_slide in result.slides:
                for shape in result_slide.shapes:
                    if not shape.has_table:
                        continue
                    table_slides += 1
                    bottom = shape.top + sum(row.height for row in shape.table.rows)
                    self.assertLessEqual(bottom, result.slide_height - Cm(1))
                    written += sum(bool(row.cells[0].text.strip())
                                   for row in list(shape.table.rows)[1:])
            self.assertGreater(table_slides, 5)
            self.assertEqual(written, 100)


if __name__ == "__main__":
    unittest.main()
